"""Report export (E-02, FR-R-06): one Excel sheet per block, or a printable PDF.

Both files are built from the same `ReportOut` the web page shows, so the numbers
always match the screen.
"""

import io
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.modules.reports.schemas import KpiOut, ReportOut

PRIMARY = "1070F8"
HEADER_FILL = PatternFill("solid", fgColor="EAF2FF")

# Vietnamese names need a font with the full Latin Extended set. Docker installs DejaVu;
# a Windows machine has Arial. Without either, accents are removed instead of printing boxes.
FONT_CANDIDATES = [
    ("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("Arial", "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
]  # fmt: skip


def _register_font() -> tuple[str, str, bool]:
    """(regular, bold, supports Vietnamese)."""
    for name, regular, bold in FONT_CANDIDATES:
        if Path(regular).exists() and Path(bold).exists():
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, regular))
                pdfmetrics.registerFont(TTFont(f"{name}-Bold", bold))
            return name, f"{name}-Bold", True
    return "Helvetica", "Helvetica-Bold", False


def _ascii(text: str) -> str:
    text = text.replace("đ", "d").replace("Đ", "D")
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()


def _percent(value: float | None) -> str:
    return "-" if value is None else f"{value * 100:.1f}%"


def _minutes(value: float | None) -> str:
    if value is None:
        return "-"
    hours, minutes = divmod(round(value), 60)
    return f"{hours}h {minutes:02d}m" if hours else f"{minutes}m"


def _kpis(report: ReportOut) -> list[tuple[str, str, str]]:
    s = report.summary

    def number(kpi: KpiOut) -> tuple[str, str]:
        fmt = lambda v: "-" if v is None else f"{v:,.0f}"  # noqa: E731
        return fmt(kpi.value), fmt(kpi.previous)

    return [
        ("Total orders", *number(s.total_orders)),
        ("Delivered", *number(s.delivered)),
        ("Success rate", _percent(s.success_rate.value), _percent(s.success_rate.previous)),
        ("On-time rate", _percent(s.on_time_rate.value), _percent(s.on_time_rate.previous)),
        ("Failed attempt rate", _percent(s.failed_attempt_rate.value),
         _percent(s.failed_attempt_rate.previous)),
        ("Avg delivery time", _minutes(s.avg_delivery_minutes.value),
         _minutes(s.avg_delivery_minutes.previous)),
        ("Incidents", *number(s.incidents)),
        ("Active drivers", f"{s.active_drivers.value or 0:.0f} / {s.total_drivers}",
         f"{s.active_drivers.previous or 0:.0f}"),
    ]  # fmt: skip


def _scope_lines(report: ReportOut) -> list[str]:
    scope = report.summary.scope
    hubs = ", ".join(scope.hub_names) if scope.hub_ids else "All hubs"
    return [
        f"Period: {scope.date_from:%d %b %Y} - {scope.date_to:%d %b %Y}",
        f"Compared with: {scope.previous_from:%d %b %Y} - {scope.previous_to:%d %b %Y}",
        f"Hubs: {hubs}",
    ]


# =============================== Excel ===============================


def _sheet(workbook: Workbook, title: str, header: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color=PRIMARY)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append(row)
    for index, column in enumerate(header, start=1):
        width = max([len(str(column))] + [len(str(row[index - 1])) for row in rows] + [8])
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 48)
    sheet.freeze_panes = "A2"


def to_xlsx(report: ReportOut) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Summary"
    summary.append(["H4T Xpress delivery report"])
    summary["A1"].font = Font(bold=True, size=14, color=PRIMARY)
    for line in _scope_lines(report):
        summary.append([line])
    summary.append([])
    summary.append(["Indicator", "This period", "Previous period"])
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True, color=PRIMARY)
        cell.fill = HEADER_FILL
    for name, now, before in _kpis(report):
        summary.append([name, now, before])
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 18

    _sheet(workbook, "Order volume", ["Period", "Orders", "Delivered"],
           [[p.label, p.orders, p.delivered] for p in report.volume])  # fmt: skip
    _sheet(workbook, "Failure reasons", ["Reason", "Failed attempts", "Share"],
           [[r.label, r.count, r.share] for r in report.failure_reasons])  # fmt: skip
    _sheet(workbook, "Incident types", ["Type", "Incidents", "Share"],
           [[r.label, r.count, r.share] for r in report.incident_types])  # fmt: skip
    _sheet(
        workbook,
        "On-time by area",
        ["Area", "Delivered", "On time", "On-time rate"],
        [[a.area, a.delivered, a.on_time, a.on_time_rate] for a in report.on_time_by_area],
    )
    _sheet(
        workbook,
        "Top drivers",
        ["Rank", "Previous rank", "Driver", "Hub", "Delivered", "Success rate", "On-time rate"],
        [[d.rank, d.previous_rank, d.name, d.hub, d.delivered, d.success_rate, d.on_time_rate]
         for d in report.top_drivers],
    )  # fmt: skip
    _sheet(
        workbook,
        "Top hubs",
        ["Rank", "Previous rank", "Hub", "Orders", "Delivered", "On-time rate"],
        [[h.rank, h.previous_rank, h.name, h.orders, h.delivered, h.on_time_rate]
         for h in report.top_hubs],
    )  # fmt: skip
    # Rates are stored as numbers and shown as percentages, so they stay usable in Excel.
    for sheet in workbook.worksheets[1:]:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                header = sheet.cell(row=1, column=cell.column).value
                if isinstance(header, str) and ("rate" in header or header == "Share"):
                    cell.number_format = "0.0%"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# =============================== PDF ===============================


def to_pdf(report: ReportOut) -> bytes:
    regular, bold, unicode_ok = _register_font()
    text = (lambda value: str(value)) if unicode_ok else (lambda value: _ascii(str(value)))
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontName=regular, fontSize=9, leading=12)
    title = ParagraphStyle("title", parent=body, fontName=bold, fontSize=16, leading=20,
                           textColor=colors.HexColor(f"#{PRIMARY}"))  # fmt: skip
    heading = ParagraphStyle("heading", parent=body, fontName=bold, fontSize=11, leading=14,
                             spaceBefore=8, spaceAfter=4)  # fmt: skip

    def table(
        header: list[str], rows: list[list[str]], widths: list[float], text_columns: int = 1
    ) -> Table:
        """Numbers are right-aligned; the first `text_columns` columns hold text."""
        data = [header, *[[text(cell) for cell in row] for row in rows]]
        result = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
        result.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, 0), bold),
                    ("FONTNAME", (0, 1), (-1, -1), regular),
                    ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(f"#{PRIMARY}")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF2FF")),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E6E7EE")),
                    ("ALIGN", (text_columns, 0), (-1, -1), "RIGHT"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        return result

    story: list[object] = [Paragraph("H4T Xpress delivery report", title)]
    story += [Paragraph(text(line), body) for line in _scope_lines(report)]
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("Key indicators", heading))
    story.append(table(["Indicator", "This period", "Previous"],
                       [list(row) for row in _kpis(report)], [70, 45, 45]))  # fmt: skip

    story.append(Paragraph("Order volume", heading))
    chart = Drawing(170 * mm, 60 * mm)
    bars = VerticalBarChart()
    bars.x, bars.y, bars.width, bars.height = 10 * mm, 8 * mm, 155 * mm, 48 * mm
    bars.data = [[p.orders for p in report.volume], [p.delivered for p in report.volume]]
    bars.bars[0].fillColor = colors.HexColor("#A9C9FD")
    bars.bars[1].fillColor = colors.HexColor(f"#{PRIMARY}")
    bars.bars.strokeColor = None
    step = max(1, len(report.volume) // 10)
    bars.categoryAxis.categoryNames = [
        p.label if index % step == 0 else "" for index, p in enumerate(report.volume)
    ]
    bars.categoryAxis.labels.fontName = regular
    bars.categoryAxis.labels.fontSize = 6
    bars.valueAxis.labels.fontName = regular
    bars.valueAxis.labels.fontSize = 6
    bars.valueAxis.valueMin = 0
    chart.add(bars)
    story.append(chart)
    story.append(Paragraph("Light blue: orders created. Blue: delivered.", body))

    story.append(Paragraph("Top failure reasons", heading))
    story.append(table(["Reason", "Failed attempts", "Share"],
                       [[r.label, str(r.count), _percent(r.share)] for r in report.failure_reasons],
                       [80, 40, 40]))  # fmt: skip
    story.append(Paragraph("Incidents by type", heading))
    story.append(table(["Type", "Incidents", "Share"],
                       [[r.label, str(r.count), _percent(r.share)] for r in report.incident_types],
                       [80, 40, 40]))  # fmt: skip
    story.append(Paragraph("On-time delivery by area", heading))
    story.append(table(["Area", "Delivered", "On time", "Rate"],
                       [[a.area, str(a.delivered), str(a.on_time), _percent(a.on_time_rate)]
                        for a in report.on_time_by_area], [70, 30, 30, 30]))  # fmt: skip
    story.append(Paragraph("Top drivers", heading))
    story.append(table(["#", "Driver", "Hub", "Delivered", "Success", "On time"],
                       [[str(d.rank), d.name, d.hub, str(d.delivered), _percent(d.success_rate),
                         _percent(d.on_time_rate)] for d in report.top_drivers],
                       [10, 50, 40, 25, 25, 25], text_columns=3))  # fmt: skip
    story.append(Paragraph("Top hubs", heading))
    story.append(table(["#", "Hub", "Orders", "Delivered", "On time"],
                       [[str(h.rank), h.name, str(h.orders), str(h.delivered),
                         _percent(h.on_time_rate)] for h in report.top_hubs],
                       [10, 60, 30, 30, 30], text_columns=2))  # fmt: skip

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm, title="H4T Xpress delivery report",
    )  # fmt: skip
    document.build(story)
    return buffer.getvalue()
