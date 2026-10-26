"""Reports (E-01, E-02; FR-R-01..FR-R-06): numbers match the data, hub scope is enforced,
and both export formats open."""

import io
from datetime import timedelta

from httpx import AsyncClient
from openpyxl import load_workbook

from app.core.clock import local_today
from tests.helpers import EXECUTIVE, HUB_MANAGER, error_code, login, scalar

COUNT_ORDERS = """
select count(*) from orders
where created_at >= (cast(:start as date)::timestamp at time zone 'Asia/Ho_Chi_Minh')
  and created_at < (cast(:end as date)::timestamp at time zone 'Asia/Ho_Chi_Minh')
"""


async def test_executive_sees_every_hub_and_numbers_match(client: AsyncClient) -> None:
    executive = await login(client, EXECUTIVE)
    response = await client.get("/api/v1/reports", headers=executive)
    assert response.status_code == 200, response.text
    report = response.json()
    scope = report["summary"]["scope"]
    assert scope["hub_ids"] == [] and len(scope["hub_names"]) == 3
    assert report["granularity"] == "day" and len(report["volume"]) == 30

    today = local_today()
    expected = await scalar(
        COUNT_ORDERS, start=today - timedelta(days=29), end=today + timedelta(days=1)
    )
    assert report["summary"]["total_orders"]["value"] == expected
    assert sum(point["orders"] for point in report["volume"]) == expected
    # The per-day split adds up to the period KPIs.
    delivered = sum(point["delivered"] for point in report["volume"])
    assert delivered == report["summary"]["delivered"]["value"]
    assert all(point["on_time"] <= point["delivered"] for point in report["volume"])
    rate = report["summary"]["success_rate"]["value"]
    assert rate is None or 0 <= rate <= 1
    assert len(report["top_hubs"]) <= 3
    shares = [row["share"] for row in report["failure_reasons"]]
    assert not shares or abs(sum(shares) - 1) < 0.01


async def test_hub_manager_is_limited_to_their_hub(client: AsyncClient) -> None:
    manager = await login(client, HUB_MANAGER)
    report = (await client.get("/api/v1/reports", headers=manager)).json()
    own = report["summary"]["scope"]["hub_ids"]
    assert len(own) == 1
    assert all(
        row["hub"] == report["summary"]["scope"]["hub_names"][0] for row in report["top_drivers"]
    )

    other = await scalar("select id from hubs where id <> :id limit 1", id=own[0])
    response = await client.get(
        "/api/v1/reports", params={"hub_ids": [str(other)]}, headers=manager
    )
    assert response.status_code == 403 and error_code(response) == "OUT_OF_SCOPE"

    export = await client.get("/api/v1/reports/export", headers=manager)
    assert export.status_code == 403  # report.export is an executive permission


async def test_period_and_granularity(client: AsyncClient) -> None:
    executive = await login(client, EXECUTIVE)
    today = local_today()
    bad = await client.get(
        "/api/v1/reports",
        params={"date_from": today.isoformat(), "date_to": (today - timedelta(days=3)).isoformat()},
        headers=executive,
    )
    assert error_code(bad) == "DATE_RANGE"
    monthly = (
        await client.get(
            "/api/v1/reports",
            params={
                "date_from": (today - timedelta(days=89)).isoformat(),
                "date_to": today.isoformat(),
                "granularity": "month",
            },
            headers=executive,
        )
    ).json()
    assert 3 <= len(monthly["volume"]) <= 4
    assert monthly["summary"]["scope"]["previous_to"] == (today - timedelta(days=90)).isoformat()


async def test_exports_open_as_excel_and_pdf(client: AsyncClient) -> None:
    executive = await login(client, EXECUTIVE)
    xlsx = await client.get("/api/v1/reports/export", params={"format": "xlsx"}, headers=executive)
    assert xlsx.status_code == 200
    assert "attachment" in xlsx.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames == [
        "Summary",
        "Order volume",
        "Failure reasons",
        "Incident types",
        "On-time by area",
        "Top drivers",
        "Top hubs",
    ]
    assert workbook["Order volume"].max_row == 31  # header + 30 days

    pdf = await client.get("/api/v1/reports/export", params={"format": "pdf"}, headers=executive)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert pdf.headers["content-type"] == "application/pdf"
    audited = await scalar("select count(*) from audit_logs where action = 'report.exported'")
    assert audited >= 2
